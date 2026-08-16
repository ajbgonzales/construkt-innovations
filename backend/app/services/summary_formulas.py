from re import search

from openpyxl.utils import get_column_letter

from services.utils import get_header_col_letter


def _get_start_and_last_date_col_letters(ws):
    header_row = 1
    date_columns = [
        cell.column
        for cell in ws[header_row]
        if search(r"\d{4}-\d{2}-\d{2}", cell.value)
    ]

    start_col = min(date_columns)
    last_col = max(date_columns)

    return get_column_letter(start_col), get_column_letter(last_col)


def get_manpower(ws):
    start_col_letter, last_col_letter = _get_start_and_last_date_col_letters(ws)
    # Sum of all work hours / 8 / 6
    return f"=ROUND(SUM('{ws.title}'!{start_col_letter}2:{last_col_letter}{ws.max_row})/48,2)"


def get_net_amount(ws):
    col_letter = get_header_col_letter(ws, "Net Amount")
    return f"=SUM('{ws.title}'!{col_letter}2:{col_letter}{ws.max_row})"


PROJECT_TOTALS_EXCLUDED_COLUMNS = {
    "Employee ID",
    "Employee Full Name",
    "Position",
    "Project",
    "Is Flagged",
    "Notes",
}


def get_project_totals_row(ws):
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    last_row = ws.max_row

    excluded_cols = [
        col_num
        for col_num, header in enumerate(header_row, start=1)
        if header in PROJECT_TOTALS_EXCLUDED_COLUMNS
    ]
    label_col = min(excluded_cols)

    totals_row = []
    for col_num, header in enumerate(header_row, start=1):
        if header in PROJECT_TOTALS_EXCLUDED_COLUMNS:
            totals_row.append("TOTAL" if col_num == label_col else None)
        else:
            col_letter = get_column_letter(col_num)
            totals_row.append(f"=SUM({col_letter}2:{col_letter}{last_row})")

    ws.append(totals_row)
    totals_row_num = ws.max_row
    ws.merge_cells(
        start_row=totals_row_num,
        start_column=label_col,
        end_row=totals_row_num,
        end_column=max(excluded_cols),
    )


def get_total_disbursement(ws):
    manpower_col_letter = get_header_col_letter(ws, "Manpower")
    net_amount_col_letter = get_header_col_letter(ws, "Net Amount")
    ws.append(
        [
            "TOTAL Disbursement",
            f"=SUM({manpower_col_letter}{2}:{manpower_col_letter}{ws.max_row})",
            f"=SUM({net_amount_col_letter}{2}:{net_amount_col_letter}{ws.max_row})",
        ]
    )


def format_number_cells(ws, column_names):
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    for name in column_names:
        col_num = header_row.index(name) + 1
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col_num).number_format = "#,##0.00"
