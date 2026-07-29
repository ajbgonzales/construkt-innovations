import re

from openpyxl.utils import get_column_letter
from pandas import DataFrame

from services.dataframe import get_loc_given_substring
from services.dates import get_week_range


def get_employee_attribute(df: DataFrame, rows: list, index: int, label: str):
    _, col = get_loc_given_substring(df, label)
    cols = df.columns.tolist()
    name_col = cols[cols.index(col) + 1]
    return getattr(rows[index], name_col)


def get_header_col_letter(ws, header_name):
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col_num = header_row.index(header_name) + 1
    return get_column_letter(col_num)


def get_header_by_regex(header_row: tuple, pattern: str) -> str | None:
    return next(
        (
            value
            for value in header_row
            if value is not None and re.search(pattern, str(value))
        ),
        None,
    )


def get_work_week_dates(ws):
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    header_date = get_header_by_regex(header_row, r"\d{4}-\d{2}-\d{2}")
    start, end = get_week_range(header_date)
    return [start, end]


def generate_filename(start_date, end_date):
    return f"Daily Payroll from {start_date.strftime('%B %d')} to {end_date.strftime('%B %d %Y')}"
